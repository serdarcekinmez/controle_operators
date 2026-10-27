"""
Parcours utilisateur complet contre l'exécutable (ou le code source), par HTTP.

    python tests/smoke_journey.py create <exe|source> <dossier_de_données> <etat.json>
    python tests/smoke_journey.py verify <exe|source> <dossier_de_données> <etat.json>
    python tests/smoke_journey.py mismatch <exe_ancien> <exe_nouveau> <dossier_de_données>

create : lance, vérifie l'interface, crée un contrôle avec documents PDF et
image, génère PDF et fichier Excel (création, mise à jour, suppression,
réouverture), quitte pendant qu'un enregistrement est en cours.
verify : relance (éventuellement un exécutable remplacé par une version plus
récente) et retrouve rapports, documents, historique et configuration.
mismatch : version ancienne ouverte + lancement de la nouvelle -> message clair.
Le dossier de données doit être temporaire (jamais celui d'un utilisateur).
"""
import base64
import io
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OPENER = urllib.request.build_opener(urllib.request.ProxyHandler({}))
failures = []


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + (f" — {detail}" if detail and not condition else ""), flush=True)
    if not condition:
        failures.append(name)
    return condition


def command(target):
    return [sys.executable, str(ROOT / "server.py")] if target == "source" else [str(Path(target).resolve())]


def env_for(data):
    return dict(os.environ, CONTROLE_AGENCES_DATA_DIR=str(data), CONTROLE_AGENCES_NO_DIALOG="1",
                PYTHONIOENCODING="utf-8")


def call(port, path, method="GET", body=None, headers=None, timeout=120):
    request = urllib.request.Request(f"http://127.0.0.1:{port}{path}", data=body, method=method, headers=headers or {})
    try:
        with OPENER.open(request, timeout=timeout) as r:
            return r.status, r.read(), r.headers.get("Content-Type", "")
    except urllib.error.HTTPError as e:
        return e.code, e.read(), e.headers.get("Content-Type", "")
    except OSError:
        return None, b"", ""


def post_json(port, path, payload):
    status, data, _ = call(port, path, "POST", json.dumps(payload).encode(), {"Content-Type": "application/json"})
    return status, (json.loads(data) if data else {})


def post_form(port, path, fields, files):
    boundary = uuid.uuid4().hex
    parts = []
    for name, value in fields.items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode())
    for name, filename, content in files:
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'
                     f'Content-Type: application/octet-stream\r\n\r\n'.encode("utf-8") + content + b"\r\n")
    body = b"".join(parts) + f"--{boundary}--\r\n".encode()
    status, data, _ = call(port, path, "POST", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return status, (json.loads(data) if data else {})


def pdf_bytes(pages):
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    for i in range(pages):
        c.drawString(100, 700, f"Annexe page {i + 1}")
        c.showPage()
    c.save()
    return buffer.getvalue()


def image_bytes(fmt, size=(640, 480), noise=False):
    buffer = io.BytesIO()
    image = Image.frombytes("RGB", size, os.urandom(size[0] * size[1] * 3)) if noise else Image.new("RGB", size, "navy")
    image.save(buffer, format=fmt, quality=95)
    return buffer.getvalue()


def start(target, data, wait=60):
    process = subprocess.Popen(command(target) + ["--no-browser"], cwd=ROOT, env=env_for(data),
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    end = time.monotonic() + wait
    while time.monotonic() < end:
        try:
            info = json.loads((Path(data) / "instance.json").read_text(encoding="utf-8"))
            status, body, _ = call(info["port"], "/api/identity", timeout=2)
            if status == 200 and json.loads(body).get("instance") == info["instance"]:
                return process, info["port"], json.loads(body)
        except (OSError, ValueError, KeyError):
            pass
        if process.poll() is not None:
            break
        time.sleep(0.2)
    process.kill()
    raise SystemExit(f"L'application n'a pas démarré (code {process.poll()}).")


def stop(process, port):
    post_json(port, "/api/quit", {})
    try:
        return process.wait(60)
    except subprocess.TimeoutExpired:
        process.kill()
        return None


REPORT = dict(branch="Gabriel Péri", controller="Zoë Brûlé", control_date="2026-10-27",
              control_time="09:15", statuses={"controle_caisses": "ok", "controle_cash": "problem"},
              problem_comments={"controle_cash": "Écart de 5 €"}, observations="Contrôle de fumée V5")


def create(target, data, state_file):
    data = Path(data)
    data.mkdir(parents=True, exist_ok=True)
    (data / "config.json").write_text(json.dumps({"default_controller_name": "Zoë Brûlé"}, ensure_ascii=False),
                                      encoding="utf-8")
    process, port, ident = start(target, data)
    print(f"Démarré : port {port}, version {ident.get('version')}")

    # Interface React complète.
    status, html, _ = call(port, "/")
    check("interface : index.html", status == 200 and b'id="root"' in html)
    assets = re.findall(r'(?:src|href)="/([^"]+)"', html.decode("utf-8"))
    check("interface : fichiers référencés servis",
          assets and all(call(port, "/" + a)[0] == 200 for a in assets), assets)
    status, body, _ = call(port, "/api/config")
    check("configuration lue", status == 200 and json.loads(body)["default_controller"] == "Zoë Brûlé")

    # Vérification des documents à l'ajout.
    st, res = post_form(port, "/api/attachments/check", {}, [("file", "annexe.pdf", pdf_bytes(2))])
    check("document PDF accepté", st == 200 and res.get("ok") and res.get("pages") == 2, res)
    st, res = post_form(port, "/api/attachments/check", {}, [("file", "photo.jpg", image_bytes("JPEG"))])
    check("image acceptée", st == 200 and res.get("ok"), res)

    # Contrôle avec documents -> PDF final.
    st, res = post_form(port, "/api/report", {"data": json.dumps(REPORT, ensure_ascii=False)}, [
        ("files", "annexe.pdf", pdf_bytes(2)), ("files", "scan.png", image_bytes("PNG")),
        ("files", "photo é.jpg", image_bytes("JPEG"))])
    check("PDF généré", st == 200 and res.get("url"), res)
    report_id = res["report"]["report_id"]
    status, content, ctype = call(port, res["url"])
    pages = len(PdfReader(io.BytesIO(content)).pages) if status == 200 else 0
    check("PDF : rapport + 2 pages PDF + 2 images", status == 200 and "pdf" in ctype and pages >= 5, f"{pages} pages")
    check("PDF : nom avec l'heure", "_09h15_" in res.get("filename", ""), res.get("filename"))

    # Second contrôle le même jour, même contrôleur : rien n'est écrasé.
    st, res2 = post_form(port, "/api/report", {"data": json.dumps(dict(REPORT, control_time="15:40"), ensure_ascii=False)}, [])
    check("second PDF du jour distinct", st == 200 and res2.get("filename") != res.get("filename"), res2.get("filename"))
    status, again, _ = call(port, res["url"])
    check("premier PDF intact", again == content)

    # Fichier de contrôle Excel : création, ajout, mise à jour, suppression, réouverture.
    st, s0 = post_json(port, "/api/excel/state", {"create": True})
    check("Excel : nouveau fichier", st == 200 and s0["workbook"] and s0["total"] == 0)
    report_row = dict(REPORT, report_id=report_id, filename=res["filename"])
    st, s1 = post_json(port, "/api/excel/save", {"base": s0["workbook"], "session_rows": [], "deleted": [], "report": report_row})
    check("Excel : contrôle ajouté", st == 200 and s1["total"] == 1, s1)
    st, s2 = post_json(port, "/api/excel/save", {"base": s1["workbook"], "session_rows": [], "deleted": [],
                                                 "report": dict(report_row, observations="corrigé")})
    check("Excel : mise à jour sans doublon", st == 200 and s2["total"] == 1)
    other = dict(REPORT, controller="Ana", branch="Lyon", control_time="10:00")
    st, s3 = post_json(port, "/api/excel/save", {"base": s2["workbook"], "session_rows": [], "deleted": [], "report": other})
    check("Excel : second contrôle", st == 200 and s3["total"] == 2)
    st, s4 = post_json(port, "/api/excel/delete", {"base": s3["workbook"], "session_rows": [], "deleted": [],
                                                   "controller": "ana", "date": "2026-10-27", "time": "10:00"})
    check("Excel : suppression", st == 200 and s4["total"] == 1)
    workbook = data.parent / "Contrôle agences.xlsx"   # fichier géré par l'utilisateur
    workbook.write_bytes(base64.b64decode(s4["workbook"]))
    sheet = load_workbook(workbook).worksheets[0]
    check("Excel : ouvrable par openpyxl", sheet.max_row >= 2)

    # Quitter pendant un enregistrement : le PDF en cours est terminé.
    heavy = [("files", f"grande_{i}.jpg", image_bytes("JPEG", (2600, 2000), noise=True)) for i in range(6)]
    outcome = {}

    def slow_report():
        outcome["result"] = post_form(port, "/api/report",
                                      {"data": json.dumps(dict(REPORT, control_time="17:05"), ensure_ascii=False)}, heavy)

    worker = threading.Thread(target=slow_report)
    worker.start()
    time.sleep(1.0)
    code = stop(process, port)
    worker.join(120)
    st, res3 = outcome.get("result", (None, {}))
    check("Quitter pendant un enregistrement : la requête se termine", st == 200, (st, res3))
    check("Quitter : le processus s'arrête (code 0)", code == 0, code)

    Path(state_file).write_text(json.dumps({
        "report_ids": [report_id, res2["report"]["report_id"]] + ([res3["report"]["report_id"]] if st == 200 else []),
        "pdf": base64.b64encode(content).decode(), "url": res["url"], "workbook": str(workbook),
        "version": ident.get("version")}), encoding="utf-8")


def verify(target, data, state_file):
    state = json.loads(Path(state_file).read_text(encoding="utf-8"))
    process, port, ident = start(target, data)
    print(f"Relancé : port {port}, version {ident.get('version')} (avant : {state['version']})")
    status, body, _ = call(port, "/api/history")
    items = {i["report_id"]: i for i in json.loads(body)["items"]}
    check("historique : tous les contrôles retrouvés", all(r in items for r in state["report_ids"]),
          [r for r in state["report_ids"] if r not in items])
    check("historique : PDF présents", all(items.get(r, {}).get("has_pdf") for r in state["report_ids"]))
    status, content, _ = call(port, state["url"])
    check("PDF identique après relance", status == 200 and content == base64.b64decode(state["pdf"]))
    for rid in state["report_ids"][2:]:
        status, content, _ = call(port, f"/api/reports/{rid}")
        check("PDF enregistré pendant la fermeture : complet",
              status == 200 and len(PdfReader(io.BytesIO(content)).pages) >= 7)
    status, body, _ = call(port, "/api/config")
    check("configuration conservée", json.loads(body)["default_controller"] == "Zoë Brûlé")
    uploads = list((Path(data) / "uploads").rglob("*.*"))
    check("documents joints conservés", len([u for u in uploads if not u.name.startswith(".")]) >= 3, len(uploads))
    check("aucun fichier temporaire restant", not list(Path(data).rglob("*.part")))
    reopened = base64.b64encode(Path(state["workbook"]).read_bytes()).decode()
    st, s = post_json(port, "/api/excel/state", {"base": reopened})
    check("Excel réouvert : contrôle conservé", st == 200 and s["total"] == 1)
    check("Quitter : code 0", stop(process, port) == 0)


def mismatch(old, new, data):
    process, port, ident = start(old, data)
    result = subprocess.run(command(new) + ["--no-browser"], cwd=ROOT, env=env_for(data), timeout=90)
    log = (Path(data) / "logs" / "app.log").read_text(encoding="utf-8", errors="replace")
    check("autre version déjà ouverte : refus (code 1)", result.returncode == 1, result.returncode)
    check("autre version déjà ouverte : message clair", "Une autre version de Contrôle Agences" in log)
    check("aucun second serveur", all(call(p, "/api/identity", timeout=2)[0] != 200 for p in (8600, 8601, 8602) if p != port))
    stop(process, port)


if __name__ == "__main__":
    action = sys.argv[1]
    {"create": create, "verify": verify, "mismatch": mismatch}[action](*sys.argv[2:5])
    if failures:
        print(f"ÉCHEC : {failures}")
        sys.exit(1)
    print(f"Parcours « {action} » : toutes les vérifications sont passées.")
