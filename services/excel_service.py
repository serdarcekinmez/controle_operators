"""
Service Excel : fichier de contrôle local du contrôleur (.xlsx).

Chaque contrôleur conserve SON fichier Excel sur son poste. À chaque
contrôle, il le téléverse dans l'application, qui y ajoute la ligne du
contrôle (ou la met à jour si le même contrôle y figure déjà), trie les
lignes par date puis par heure, et le lui rend à télécharger. Un contrôle
peut aussi en être retiré (même date, même heure, même contrôleur).

Aucun compte en ligne, aucun envoi réseau : tout se fait en mémoire.

L'onglet « Controles » contient tous les contrôles ; un onglet par agence
(« Lyon », « Rivoli »...) reprend ceux de l'agence. Les onglets d'agence
sont recréés à chaque mise à jour ; les autres onglets du classeur et les
colonnes ajoutées à la main dans « Controles » sont conservés.

Les fonctions lèvent ExcelError avec un message déjà rédigé pour
l'utilisateur final.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from constants import (
    BRANCHES,
    CONTROL_KEY_TO_SHEET_COLUMN,
    EXCEL_SHEET_NAME,
    SHEET_COLUMN_LABELS,
    SHEET_COLUMNS,
    SHEET_LEGACY_TIMESTAMP_LABELS,
    SHEET_PROBLEM_PREFIXES,
    SHEET_VALUE_OK,
    STATUS_PROBLEM,
    sheet_cell_value,
)
from services.logging_service import get_logger

logger = get_logger()


class ExcelError(Exception):
    """Erreur de lecture / écriture du fichier Excel, message utilisateur."""


# Colonnes au format date / heure / nombre dans le classeur (le reste est du
# texte).
_DATE_COLUMNS = {"date_controle": "DD/MM/YYYY", "heure_controle": "HH:MM"}
_INT_COLUMNS = {"nb_problemes"}
_STATUS_COLUMNS = set(CONTROL_KEY_TO_SHEET_COLUMN.values())

# Clé interne de l'ancienne colonne « Enregistré le » (voir constants).
_LEGACY_TIMESTAMP = "horodatage"

_HEADER_FILL = PatternFill("solid", fgColor="9D174D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_OK_FONT = Font(bold=True, color="1A7F37")
_PB_FONT = Font(bold=True, color="B42318")
_PB_FILL = PatternFill("solid", fgColor="FDECEA")

_COLUMN_WIDTHS = {
    "report_id": 14, "date_controle": 12, "heure_controle": 8, "agence": 16,
    "controleur": 20, "nb_problemes": 11, "observations": 45,
    "fichier_pdf": 38, "email_envoye": 9,
}
_STATUS_WIDTH = 28


# --------------------------------------------------------------------------- #
# Construction d'une ligne
# --------------------------------------------------------------------------- #
def build_row(report: dict) -> dict[str, str]:
    """
    Ligne Excel d'un contrôle, à partir du rapport généré dans l'application.

    Chaque point de contrôle devient une colonne contenant « RAS »,
    « ANOMALIE — commentaire » ou « Non renseigné ».
    """
    statuses = report.get("statuses") or {}
    comments = report.get("problem_comments") or {}
    row: dict[str, str] = {
        "report_id": report.get("report_id", ""),
        "date_controle": report.get("control_date", ""),
        "heure_controle": _to_time(report.get("control_time", "")),
        "agence": report.get("branch", ""),
        "controleur": report.get("controller", ""),
        "observations": (report.get("observations") or "").strip(),
        "fichier_pdf": report.get("filename", ""),
        "email_envoye": "Oui" if report.get("email_sent") else "Non",
    }
    problems = 0
    for key, column in CONTROL_KEY_TO_SHEET_COLUMN.items():
        status = statuses.get(key) or ""
        row[column] = sheet_cell_value(status, comments.get(key))
        if status == STATUS_PROBLEM:
            problems += 1
    row["nb_problemes"] = str(problems)
    return row


def normalize_name(name: str | None) -> str:
    """Nom comparable : espaces superflus retirés, casse ignorée."""
    return " ".join((name or "").split()).casefold()


def row_key(row: dict[str, str]) -> tuple[str, str, str]:
    """
    Identité d'un contrôle : (date ISO, heure HH:MM, contrôleur normalisé).

    Enregistrer deux fois le même contrôle met à jour sa ligne ; c'est aussi
    ce que l'on saisit pour supprimer un contrôle du fichier.
    """
    return (
        row.get("date_controle", ""),
        row.get("heure_controle", ""),
        normalize_name(row.get("controleur")),
    )


# --------------------------------------------------------------------------- #
# Normalisation des valeurs lues
# --------------------------------------------------------------------------- #
_FR_DATE_RE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?")
_TIME_RE = re.compile(r"(\d{1,2})\s*[:hH]\s*(\d{2})")


def _to_iso(value, *, with_time: bool) -> str:
    """Date Excel / texte ISO / texte JJ/MM/AAAA -> texte ISO (tri, filtres)."""
    if value is None or value == "":
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="minutes") if with_time else value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value).strip()
    match = _FR_DATE_RE.match(text)
    if match:
        day, month, year, hour, minute = match.groups()
        iso = f"{year}-{int(month):02d}-{int(day):02d}"
        if with_time and hour:
            iso += f" {int(hour):02d}:{minute}"
        return iso
    text = text.replace("T", " ")
    return text[:16] if with_time else text[:10]


def _to_time(value) -> str:
    """Heure Excel / texte (« 9:05 », « 14h30 », date + heure) -> « HH:MM »."""
    if value is None or value == "":
        return ""
    if isinstance(value, (dt.time, dt.datetime)):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)) and 0 <= value < 1:
        # Heure saisie dans une cellule sans format : fraction de journée.
        minutes = min(round(value * 24 * 60), 24 * 60 - 1)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    text = str(value).strip()
    match = _TIME_RE.search(text)
    if match and int(match.group(1)) < 24 and int(match.group(2)) < 60:
        return f"{int(match.group(1)):02d}:{match.group(2)}"
    return text


def _to_text(column: str, value) -> str:
    """Valeur de cellule -> texte, dans le format attendu par l'application."""
    if column == "date_controle":
        return _to_iso(value, with_time=False)
    if column in ("heure_controle", _LEGACY_TIMESTAMP):
        return _to_time(value)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _to_cell(column: str, text: str):
    """Texte de l'application -> valeur typée pour Excel (dates, nombres)."""
    if not text:
        return None
    if column == "date_controle":
        try:
            return dt.date.fromisoformat(text)
        except ValueError:
            return text
    if column == "heure_controle":
        try:
            return dt.time.fromisoformat(text)
        except ValueError:
            return text
    if column in _INT_COLUMNS:
        try:
            return int(text)
        except ValueError:
            return text
    return text


# --------------------------------------------------------------------------- #
# Lecture du classeur
# --------------------------------------------------------------------------- #
def _header_key(value) -> str | None:
    """Colonne de l'application correspondant à un titre d'en-tête."""
    text = str(value or "").strip().casefold()
    if not text:
        return None
    for column in SHEET_COLUMNS:
        if text in (column.casefold(), SHEET_COLUMN_LABELS[column].casefold()):
            return column
    if text in (label.casefold() for label in SHEET_LEGACY_TIMESTAMP_LABELS):
        return _LEGACY_TIMESTAMP
    return None


def _open_workbook(data: bytes):
    try:
        return load_workbook(io.BytesIO(data))
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        logger.warning("Fichier Excel illisible : %s", exc)
        raise ExcelError(
            "Ce fichier n'a pas pu être ouvert. Vérifiez qu'il s'agit bien "
            "d'un fichier Excel (.xlsx) enregistré depuis Excel, et qu'il "
            "n'est pas protégé par un mot de passe."
        ) from exc


def _find_sheet(workbook):
    """Onglet « Controles », ou premier onglet portant les bons en-têtes."""
    if EXCEL_SHEET_NAME in workbook.sheetnames:
        return workbook[EXCEL_SHEET_NAME]
    for sheet in workbook.worksheets:
        headers = {_header_key(cell.value) for cell in sheet[1]}
        if {"report_id", "date_controle", "agence"} <= headers:
            return sheet
    raise ExcelError(
        f"Ce fichier ne contient pas d'onglet « {EXCEL_SHEET_NAME} ». "
        "Téléversez le fichier de contrôle téléchargé depuis l'application "
        "(ou retirez ce fichier et utilisez le bouton rouge en haut à "
        "droite de la page)."
    )


def _read_table(sheet) -> tuple[list, list[list]]:
    """En-têtes (valeurs brutes) et lignes de données non vides."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = list(rows[0])
    while header and header[-1] in (None, ""):
        header.pop()
    data = [
        list(row[: len(header)]) + [None] * (len(header) - len(row))
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]
    return header, data


def _normalize_table(header: list, data: list[list]) -> list[str | None]:
    """
    Met le tableau au format de l'application (modifie header et data).

    - l'ancienne colonne « Enregistré le » est retirée ; son heure remplit
      « Heure » quand celle-ci est vide ;
    - les colonnes de l'application absentes sont insérées après la colonne
      qui les précède dans SHEET_COLUMNS (à défaut, tout à droite).

    Retourne la clé de chaque colonne (None pour une colonne personnelle).
    """
    keys = [_header_key(value) for value in header]

    legacy = keys.index(_LEGACY_TIMESTAMP) if _LEGACY_TIMESTAMP in keys else None
    if legacy is not None:
        del header[legacy], keys[legacy]
        stamps = [_to_time(values.pop(legacy)) for values in data]
    else:
        stamps = [""] * len(data)

    for idx, column in enumerate(SHEET_COLUMNS):
        if column in keys:
            continue
        before = [keys.index(col) for col in SHEET_COLUMNS[:idx] if col in keys]
        insert_at = max(before) + 1 if before else len(keys)
        header.insert(insert_at, SHEET_COLUMN_LABELS[column])
        keys.insert(insert_at, column)
        for values in data:
            values.insert(insert_at, None)

    time_idx = keys.index("heure_controle")
    for values, stamp in zip(data, stamps):
        if stamp and values[time_idx] in (None, ""):
            values[time_idx] = _to_cell("heure_controle", stamp)
    return keys


def _row_texts(keys: list[str | None], values: list) -> dict[str, str]:
    """Ligne du classeur -> {colonne de l'application: texte}."""
    row = {column: "" for column in SHEET_COLUMNS}
    for key, value in zip(keys, values):
        if key:
            row[key] = _to_text(key, value)
    return row


def read_rows(data: bytes) -> list[dict[str, str]]:
    """
    Relit les contrôles d'un fichier Excel de contrôle.

    Retourne une liste de lignes {colonne: texte}, dates au format ISO.
    """
    sheet = _find_sheet(_open_workbook(data))
    header, rows = _read_table(sheet)
    keys = _normalize_table(header, rows)
    return [_row_texts(keys, values) for values in rows]


# --------------------------------------------------------------------------- #
# Écriture du classeur
# --------------------------------------------------------------------------- #
def sort_key(row: dict[str, str]) -> tuple[str, str]:
    """Ordre des contrôles : date, puis heure."""
    return (row.get("date_controle", ""), row.get("heure_controle", ""))


def _style_sheet(sheet, keys: list[str | None], n_rows: int) -> None:
    """Mise en forme : en-tête coloré, statuts en couleur, filtres, volet figé."""
    for col_idx, key in enumerate(keys, start=1):
        head = sheet.cell(row=1, column=col_idx)
        if key:
            head.fill, head.font = _HEADER_FILL, _HEADER_FONT
            head.alignment = Alignment(vertical="center")
            width = _STATUS_WIDTH if key in _STATUS_COLUMNS else _COLUMN_WIDTHS.get(key, 14)
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        for row_idx in range(2, n_rows + 2):
            cell = sheet.cell(row=row_idx, column=col_idx)
            if key in _DATE_COLUMNS:
                cell.number_format = _DATE_COLUMNS[key]
            elif key in _STATUS_COLUMNS:
                text = str(cell.value or "")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if text == SHEET_VALUE_OK:
                    cell.font, cell.fill = _OK_FONT, PatternFill()
                elif text.startswith(SHEET_PROBLEM_PREFIXES):
                    cell.font, cell.fill = _PB_FONT, _PB_FILL
                else:
                    cell.font, cell.fill = Font(), PatternFill()
            elif key == "observations":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif key == "nb_problemes":
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = _PB_FONT if cell.value not in (None, 0, "0") else Font(bold=True)
    sheet.freeze_panes = "A2"
    if keys:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{n_rows + 1}"


def _write_table(
    sheet, header: list, data: list[list], keys: list[str | None], old_width: int = 0
) -> None:
    old_rows = sheet.max_row
    for row_idx, values in enumerate([header, *data], start=1):
        for col_idx, (key, value) in enumerate(zip(keys, values), start=1):
            # Affectation explicite : sheet.cell(value=None) n'efface pas.
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.value = value
            # openpyxl lit un texte commençant par « = » comme une formule :
            # les textes saisis dans l'application restent du texte.
            if row_idx > 1 and key and isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"
    if old_rows > len(data) + 1:
        sheet.delete_rows(len(data) + 2, old_rows - len(data) - 1)
    # Colonne retirée (ancienne « Enregistré le ») : la dernière est libérée.
    if old_width > len(header):
        sheet.delete_cols(len(header) + 1, old_width - len(header))
    _style_sheet(sheet, keys, len(data))


_SHEET_TITLE_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def _branch_sheet_title(branch) -> str:
    """Nom d'onglet Excel valide pour une agence (31 caractères au plus)."""
    return _SHEET_TITLE_FORBIDDEN.sub("", str(branch or "")).strip(" '")[:31].strip()


def _write_branch_sheets(workbook, main_sheet, data: list[list], position: dict) -> None:
    """
    Un onglet par agence, après « Controles », avec les contrôles de l'agence.

    Les onglets portant le nom d'une agence sont recréés à chaque mise à
    jour (une agence sans contrôle n'a pas d'onglet).
    """
    by_branch: dict[str, list[list]] = {}
    for values in data:
        title = _branch_sheet_title(_to_text("agence", values[position["agence"]]))
        if title and title.casefold() != main_sheet.title.casefold():
            by_branch.setdefault(title, []).append(
                [values[position[column]] for column in SHEET_COLUMNS]
            )

    managed = {title.casefold() for title in by_branch}
    managed |= {_branch_sheet_title(branch).casefold() for branch in BRANCHES}
    managed.discard(main_sheet.title.casefold())
    for sheet in list(workbook.worksheets):
        if sheet.title.casefold() in managed:
            workbook.remove(sheet)

    header = [SHEET_COLUMN_LABELS[column] for column in SHEET_COLUMNS]
    index = workbook.worksheets.index(main_sheet)
    for offset, title in enumerate(sorted(by_branch, key=str.casefold), start=1):
        sheet = workbook.create_sheet(title, index + offset)
        _write_table(sheet, header, by_branch[title], list(SHEET_COLUMNS))
    workbook.active = index


def update_workbook(
    base: bytes | None,
    new_rows: list[dict[str, str]],
    delete_keys: tuple | list | set = (),
) -> bytes:
    """
    Ajoute, met à jour ou retire des contrôles dans un classeur.

    - base = None -> nouveau classeur avec l'onglet « Controles » ;
    - delete_keys : identités (row_key) des contrôles à retirer ;
    - new_rows : contrôles à ajouter ; une ligne de même « Réf. » ou de même
      identité (date, heure, contrôleur) est mise à jour, jamais dupliquée.

    L'onglet des contrôles est trié par date puis par heure ; les onglets
    d'agence sont recréés. Les colonnes et onglets ajoutés par l'utilisateur
    sont conservés.
    """
    if base:
        workbook = _open_workbook(base)
        sheet = _find_sheet(workbook)
        header, data = _read_table(sheet)
    else:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = EXCEL_SHEET_NAME
        header, data = [], []

    old_width = len(header)
    keys = _normalize_table(header, data)
    position = {key: idx for idx, key in enumerate(keys) if key}

    # Suppressions.
    delete_keys = set(delete_keys)
    texts = [_row_texts(keys, values) for values in data]
    kept = [
        (values, row) for values, row in zip(data, texts)
        if row_key(row) not in delete_keys
    ]

    # Insertion / mise à jour, par « Réf. » puis par identité.
    by_ref = {row["report_id"]: idx for idx, (_, row) in enumerate(kept) if row["report_id"]}
    by_key = {row_key(row): idx for idx, (_, row) in enumerate(kept)}
    for row in new_rows:
        idx = by_ref.get(row.get("report_id") or None, by_key.get(row_key(row)))
        if idx is None:
            idx = len(kept)
            kept.append(([None] * len(header), {}))
        values, previous = kept[idx]
        # Un contrôle déjà dans le fichier garde sa « Réf. ».
        row = dict(row, report_id=previous.get("report_id") or row.get("report_id", ""))
        for column in SHEET_COLUMNS:
            values[position[column]] = _to_cell(column, row.get(column, ""))
        kept[idx] = (values, row)
        if row.get("report_id"):
            by_ref[row["report_id"]] = idx
        by_key[row_key(row)] = idx

    # Tri par date puis heure (tri stable : à égalité, ordre d'ajout).
    kept.sort(key=lambda item: sort_key(item[1]))
    data = [values for values, _ in kept]

    _write_table(sheet, header, data, keys, old_width)
    _write_branch_sheets(workbook, sheet, data, position)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_workbook(rows: list[dict[str, str]]) -> bytes:
    """Nouveau classeur contenant ces lignes (export de la consultation)."""
    return update_workbook(None, rows)


def summarize(rows: list[dict[str, str]]) -> tuple[int, str]:
    """Nombre de contrôles et date du plus récent (texte ISO, vide si aucun)."""
    dates = [row["date_controle"] for row in rows if row.get("date_controle")]
    return len(rows), max(dates, default="")
